"""바이럴·운영 항목: 첫 컷, 도입형 제목, 이어지는 이야기, 인물 재생목록, 통독 달력, 성과 기록, 채널 재생목록."""
import os
import re
from datetime import date, timedelta

from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIRST_CUT_FILE = os.path.join(ROOT, "data", "첫컷_시즌1.lst")
FIRST_CUT_DEFAULT = "대본 작성 시 확정 (원칙: 가장 강한 장면을 먼저 1~2초)"

# 인물별 재생목록: (이름, 찾을 이름들, 구약/신약)
ARCS = [
    ("노아", ("노아",), "구약"), ("아브라함", ("아브라함", "아브람"), "구약"),
    ("야곱", ("야곱",), "구약"), ("요셉", ("요셉",), "구약"), ("모세", ("모세",), "구약"),
    ("여호수아", ("여호수아",), "구약"), ("기드온", ("기드온",), "구약"), ("삼손", ("삼손",), "구약"),
    ("룻", ("룻",), "구약"), ("사무엘", ("사무엘",), "구약"), ("사울", ("사울",), "구약"),
    ("다윗", ("다윗",), "구약"), ("솔로몬", ("솔로몬",), "구약"), ("엘리야", ("엘리야",), "구약"),
    ("엘리사", ("엘리사",), "구약"), ("에스더", ("에스더",), "구약"), ("욥", ("욥",), "구약"),
    ("다니엘", ("다니엘",), "구약"), ("요나", ("요나",), "구약"),
    ("세례 요한", ("세례 요한",), "신약"), ("베드로", ("베드로",), "신약"), ("바울", ("바울", "사울"), "신약"),
]


def load_first_cuts(rows):
    index = {(r["book"]["name"], r["rng"]): r for r in rows}
    for raw in open(FIRST_CUT_FILE, encoding="utf-8"):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        book, rng, cut = line.split("|")
        if (book, rng) not in index:
            raise SystemExit(f"첫 컷 목록에 제작표에 없는 본문: {line}")
        index[(book, rng)]["first_cut"] = cut
    missing = [r["disp"] for r in rows if r["season"] == 1 and "first_cut" not in r]
    if missing:
        raise SystemExit(f"시즌1 첫 컷 누락: {missing}")
    for r in rows:
        r.setdefault("first_cut", FIRST_CUT_DEFAULT)


def hook_type(hook):
    if re.search(r"\d", hook):
        return "숫자형"
    if hook.rstrip().endswith("?"):
        return "질문형"
    if "!" in hook:
        return "감탄형"
    return "사실·반전형"


def chapter_label(r):
    if r["book"]["name"] == "시편":
        return f"시편 {r['rng']}편"
    a, _, b = r["rng"].partition("-")
    c1 = a.split(":")[0]
    c2 = b.split(":")[0] if ":" in b else c1
    return f"{r['book']['name']} {c1}장" if c1 == c2 else f"{r['book']['name']} {c1}~{c2}장"


def names(r):
    return {n.strip() for n in re.split(r"[·,/]", r["people"]) if n.strip()}


def annotate(rows):
    """도입 유형, 이어지는 이야기(끊는 지점), 인물 재생목록."""
    for t in ("구약", "신약"):
        track = sorted((r for r in rows if r["book"]["testament"] == t), key=lambda r: r["track_no"])
        for cur, nxt in zip(track, track[1:] + [None]):
            cont = (nxt is not None and nxt["b"] is cur["b"] and cur["kind"] == "장면"
                    and nxt["kind"] == "장면" and names(cur) & names(nxt))
            cur["cliff"] = (f"이어지는 이야기 → 결정적 순간에서 끊고 다음 편 '{nxt['title']}'로 연결"
                            if cont else "")
    for r in rows:
        r["hook_type"] = hook_type(r["hook"])
        r["chapter"] = chapter_label(r)
        found = []
        for arc, keys, testament in ARCS:
            if r["book"]["testament"] != testament:
                continue
            if any(k == n or n.startswith(k + "(") for k in keys for n in names(r)):
                found.append(arc)
        r["arcs"] = ", ".join(found)


def challenge_sheet(wb, f, header, body, last, psa_n, P, start_default, end_default):
    """성경 통독 챌린지 달력: 날짜별 구약·신약·시편 본문."""
    ws = wb.create_sheet("통독챌린지")
    cols = ["일차", "날짜", "요일", "구약 (영상·함께 읽을 본문)", "신약 (영상·함께 읽을 본문)",
            "시편 시리즈", "오늘 편수", "체크"]
    header(ws, 1, cols, [6, 11, 5, 40, 40, 30, 7, 6])
    days = (end_default - start_default).days + 1 + 35  # 시작일을 늦추는 경우를 위한 여유
    endf = f"MAX(MAX('A타입_제작표'!$U$2:$U${last}),MAX('시편_시리즈'!$Q$2:$Q${psa_n}))"
    A = lambda c: f"'A타입_제작표'!${c}$2:${c}${last}"
    for k in range(1, days + 1):
        i = k + 1
        ws.cell(row=i, column=1, value=f"=IF(B{i}=\"\",\"\",{k})")
        ws.cell(row=i, column=2, value=f"=IF({P['시즌1 시작일 (월요일)']}+{k - 1}>{endf},\"\",{P['시즌1 시작일 (월요일)']}+{k - 1})")
        ws.cell(row=i, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=3, value=f"=IF(B{i}=\"\",\"\",CHOOSE(WEEKDAY(B{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\"))")
        for col, t in ((4, "구약"), (5, "신약")):
            m = f"MATCH(\"{t}|\"&TEXT(B{i},\"yyyy-mm-dd\"),{A('AQ')},0)"
            ws.cell(row=i, column=col, value=(
                f"=IF(B{i}=\"\",\"\",IFERROR(INDEX({A('J')},{m})&\"  \"&INDEX({A('L')},{m}),\"\"))"))
        mp = f"MATCH(\"시편|\"&TEXT(B{i},\"yyyy-mm-dd\"),'시편_시리즈'!$Y$2:$Y${psa_n},0)"
        ws.cell(row=i, column=6, value=(
            f"=IF(B{i}=\"\",\"\",IFERROR(\"시편 \"&INDEX('시편_시리즈'!$A$2:$A${psa_n},{mp})&\"편  \"&INDEX('시편_시리즈'!$G$2:$G${psa_n},{mp}),\"\"))"))
        ws.cell(row=i, column=7, value=f"=IF(B{i}=\"\",\"\",(D{i}<>\"\")+(E{i}<>\"\")+(F{i}<>\"\"))")
        ws.cell(row=i, column=8, value="□")
    body(ws, 2, days + 1, len(cols), center_cols=(1, 2, 3, 7, 8))
    ws.freeze_panes = "D2"
    note = days + 3
    ws.cell(row=note, column=1, value="※ 매일 올라오는 1분 영상을 보고, 같은 본문을 성경에서 직접 읽는 챌린지입니다 (#1분성경챌린지). 교회·소그룹 단체 참여용으로 배포하세요.").font = f(False, 9, "595959")
    ws.cell(row=note + 1, column=1, value="※ 시즌1은 건너뛰는 본문이 있으므로, 시즌1+시즌2+시편 시리즈를 모두 마치면 성경 전체를 읽게 됩니다. 단, 마가·누가복음 일부 장은 마태복음의 같은 사건으로 대신합니다.").font = f(False, 9, "595959")
    return days


def kpi_sheet(wb, f, header, body, last, P, n_weeks):
    ws = wb.create_sheet("성과기록")
    A = lambda c: f"'A타입_제작표'!${c}$2:${c}${last}"
    ws["A1"] = "성과 기록 (주 1회)"
    ws["A1"].font = f(True, 16, "1F3864")
    ws["A2"] = ("편별 수치는 A타입_제작표 끝의 노란 칸(조회수·시청 vs 넘김·평균 시청률·구독 증가)에 입력하면 여기서 자동 집계됩니다. "
                "목표치는 유튜브 공식 기준이 아니라 크리에이터들이 공유한 경험치이므로 채널 평균과 비교해 조정하세요.")
    ws["A2"].font = f(False, 9, "595959")
    ws.column_dimensions["A"].width = 22
    for c in "BCDEFGH":
        ws.column_dimensions[c].width = 15
    ws.column_dimensions["I"].width = 40

    ws["A4"] = "목표"
    ws["A4"].font = f(True)
    targets = [("시청 vs 넘김 (%)", 70), ("평균 시청률 (%)", 80)]
    for k, (label, v) in enumerate(targets):
        ws.cell(row=5 + k, column=1, value=label).font = f()
        c = ws.cell(row=5 + k, column=2, value=v)
        c.fill = PatternFill("solid", fgColor="FFF2CC")
        c.font = f(False, 10, "0000FF")

    ws["A8"] = "도입 유형별 평균 (어떤 첫 3초가 잘 되는지 비교)"
    ws["A8"].font = f(True)
    for j, t in enumerate(["도입 유형", "편수", "입력된 편", "평균 시청 vs 넘김(%)", "평균 시청률(%)", "평균 조회수"], 1):
        c = ws.cell(row=9, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for k, t in enumerate(["질문형", "숫자형", "감탄형", "사실·반전형"]):
        r = 10 + k
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=f"=COUNTIF({A('AH')},A{r})")
        ws.cell(row=r, column=3, value=f"=COUNTIFS({A('AH')},A{r},{A('AM')},\"<>\")")
        ws.cell(row=r, column=4, value=f"=IFERROR(AVERAGEIFS({A('AM')},{A('AH')},A{r}),\"-\")")
        ws.cell(row=r, column=5, value=f"=IFERROR(AVERAGEIFS({A('AN')},{A('AH')},A{r}),\"-\")")
        ws.cell(row=r, column=6, value=f"=IFERROR(AVERAGEIFS({A('AL')},{A('AH')},A{r}),\"-\")")
    for j, t in enumerate(["제목안", "입력된 편", "평균 조회수"], 1):
        c = ws.cell(row=15, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for k, t in enumerate(["A (회차형)", "B (도입형)"]):
        r = 16 + k
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=f"=COUNTIFS({A('AP')},LEFT(A{r},1),{A('AL')},\"<>\")")
        ws.cell(row=r, column=3, value=f"=IFERROR(AVERAGEIFS({A('AL')},{A('AP')},LEFT(A{r},1)),\"-\")")
    ws.cell(row=18, column=1, value="※ A타입_제작표 '사용한 제목안' 칸에 A 또는 B를 적으면 제목 방식별로 비교됩니다. 시즌1 첫 4주는 두 방식을 번갈아 써 보길 권합니다.").font = f(False, 9, "595959")

    hr = 20
    cols = ["주차", "시작일", "업로드 편수(본편)", "조회수 합계", "평균 시청 vs 넘김(%)", "평균 시청률(%)",
            "구독 증가 합계", "목표 달성", "메모 (잘된 편·개선점)"]
    for j, t in enumerate(cols, 1):
        c = ws.cell(row=hr, column=j, value=t)
        c.font = f(True, 10, "FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for w in range(1, n_weeks + 1):
        r = hr + w
        s = f"({P['시즌1 시작일 (월요일)']}+{7 * (w - 1)})"
        rng = f"{A('U')},\">=\"&{s},{A('U')},\"<\"&({s}+7)"
        ws.cell(row=r, column=1, value=w)
        ws.cell(row=r, column=2, value=f"={s}")
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=3, value=f"=COUNTIFS({rng})")
        ws.cell(row=r, column=4, value=f"=SUMIFS({A('AL')},{rng})")
        ws.cell(row=r, column=5, value=f"=IFERROR(AVERAGEIFS({A('AM')},{rng}),\"-\")")
        ws.cell(row=r, column=6, value=f"=IFERROR(AVERAGEIFS({A('AN')},{rng}),\"-\")")
        ws.cell(row=r, column=7, value=f"=SUMIFS({A('AO')},{rng})")
        ws.cell(row=r, column=8, value=f"=IF(ISNUMBER(E{r}),IF(E{r}>=$B$5,\"달성\",\"미달\"),\"-\")")
        ws.cell(row=r, column=9).fill = PatternFill("solid", fgColor="FFF2CC")
    body(ws, hr + 1, hr + n_weeks, len(cols), center_cols=(1, 2, 3, 4, 5, 6, 7, 8))
    ws.freeze_panes = f"A{hr + 1}"


def playlist_sheet(wb, f, header, body, books, last, psa_n, n_long, n_b, psa_long_n):
    ws = wb.create_sheet("채널_재생목록")
    A = lambda c: f"'A타입_제작표'!${c}$2:${c}${last}"
    ws["A1"] = "채널 운영과 재생목록"
    ws["A1"].font = f(True, 16, "1F3864")
    ws["A2"] = "한 채널에서 본편·시편 시리즈·롱폼을 모두 운영하고 재생목록으로 나눕니다 (구독자를 한곳에 모으기 위함)."
    ws["A2"].font = f(False, 9, "595959")
    cols = ["재생목록", "종류", "편수", "기준", "운영 메모"]
    header(ws, 4, cols, [30, 12, 8, 46, 40])
    items = [
        ("시즌1 · 구약", "시즌", f"=COUNTIFS({A('C')},1,{A('D')},\"구약\")", "A타입_제작표 시즌=1, 구약", "매일 1편"),
        ("시즌1 · 신약", "시즌", f"=COUNTIFS({A('C')},1,{A('D')},\"신약\")", "A타입_제작표 시즌=1, 신약", "월·수·금·토"),
        ("시즌2 · 구약", "시즌", f"=COUNTIFS({A('C')},2,{A('D')},\"구약\")", "A타입_제작표 시즌=2, 구약", ""),
        ("시즌2 · 신약", "시즌", f"=COUNTIFS({A('C')},2,{A('D')},\"신약\")", "A타입_제작표 시즌=2, 신약", ""),
        ("시편 시리즈", "시리즈", f"=COUNT('시편_시리즈'!$A$2:$A${psa_n})", "시편_시리즈 시트", "화·목·일"),
        ("롱폼 · 시즌별", "롱폼", f"=COUNTA('롱폼_시즌별'!$B$2:$B${n_long})", "롱폼_시즌별 시트", "프리미어 공개 권장"),
        ("롱폼 · 통합본(성경 순서)", "롱폼", f"=COUNTA('롱폼_통합본'!$B$2:$B${n_b})", "롱폼_통합본 시트", "전체 완료 후 정주행용"),
        ("롱폼 · 시편", "롱폼", f"=COUNTA('시편_롱폼'!$B$2:$B${psa_long_n})", "시편_롱폼 시트", ""),
        ("1분성경 통독 챌린지", "캠페인", f"=COUNTA({A('B')})", "통독챌린지 시트", "교회·소그룹 단체 참여 안내 고정"),
    ]
    for arc, _, _ in ARCS:
        items.append((f"인물 · {arc}", "인물", f"=COUNTIF({A('AK')},\"*{arc}*\")", "A타입_제작표 '인물 재생목록' 칸", "연속 이야기 몰아보기"))
    for bk in books:
        items.append((f"권별 · {bk['name']}", "권별", f"=COUNTIF({A('I')},\"{bk['name']}\")", "A타입_제작표 책", ""))
    for k, it in enumerate(items, 5):
        for j, v in enumerate(it, 1):
            ws.cell(row=k, column=j, value=v)
    body(ws, 5, 4 + len(items), len(cols), center_cols=(2, 3))
    ws.freeze_panes = "A5"
