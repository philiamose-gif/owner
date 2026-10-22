"""시편 시리즈(150편 전체) 데이터와 시트."""
import os

from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PSA_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "시편시리즈.psa")

BOOKS = [(1, 41, "제1권"), (42, 72, "제2권"), (73, 89, "제3권"), (90, 106, "제4권"), (107, 150, "제5권")]

LONGS = [
    (1, 10, "복 있는 사람과 아침의 기도"), (11, 20, "흔들릴 때 붙드시는 하나님"),
    (21, 30, "목자와 영광의 왕"), (31, 41, "용서와 기다림"),
    (42, 50, "목마른 영혼과 피난처"), (51, 60, "회개와 광야의 기도"),
    (61, 72, "땅끝에서 부르는 노래"), (73, 80, "아삽의 질문"), (81, 89, "다시 돌이키소서"),
    (90, 98, "영원하신 왕"), (99, 106, "거룩하신 하나님의 은택"), (107, 118, "감사와 할렐"),
    (119, 127, "말씀과 순례의 노래"), (128, 137, "깊은 곳에서 부르는 순례의 노래"),
    (138, 150, "할렐루야 피날레"),
]

# 시편 시리즈 업로드 요일: 주 3편(화·목·일)
DAYS = (1, 3, 6)


def load(rows):
    main = {int(r["rng"]): r for r in rows if r["book"]["name"] == "시편"}
    out = []
    for raw in open(PSA_FILE, encoding="utf-8"):
        line = raw.rstrip("\n")
        if not line.strip() or line.startswith("#"):
            continue
        p = line.split("|")
        if len(p) != 9:
            raise SystemExit(f"시편 시리즈 형식 오류: {line}")
        no = int(p[0])
        kind, title, hook, kv, msg, author, scene, caution = p[1:]
        if title == "=본편":
            if no not in main:
                raise SystemExit(f"시편 {no}편은 본편에 없어 재사용할 수 없습니다")
            m = main[no]
            ps = dict(no=no, kind=kind, title=m["title"], hook=m["hook"], kv=m["kv"], msg=m["msg"],
                      author=author, scene=m["place"], caution=m["caution"], reuse=m)
        else:
            if no in main:
                raise SystemExit(f"시편 {no}편은 본편에 있으므로 '=본편'으로 표시해야 합니다")
            if kv.split(":")[0] != str(no):
                raise SystemExit(f"시편 {no}편 핵심 구절 장 번호 오류: {kv}")
            ps = dict(no=no, kind=kind, title=title, hook=hook, kv=f"시 {kv}", msg=msg,
                      author=author, scene=scene, caution=caution, reuse=None)
        ps["book"] = next(name for a, b, name in BOOKS if a <= no <= b)
        ps["long"] = next(f"PL{i:02d}" for i, (a, b, _) in enumerate(LONGS, 1) if a <= no <= b)
        out.append(ps)
    nums = [x["no"] for x in out]
    if nums != list(range(1, 151)):
        missing = sorted(set(range(1, 151)) - set(nums))
        dup = sorted({n for n in nums if nums.count(n) > 1})
        raise SystemExit(f"시편 시리즈 누락 {missing} / 중복 {dup} / 순서 확인 필요")
    if {x["no"] for x in out if x["reuse"]} != set(main):
        raise SystemExit("본편 시편이 시편 시리즈에서 모두 재사용되지 않았습니다")
    return out


def sheets(wb, psalms, f, header, body, P, STATUS, A):
    n = len(psalms) + 1
    ws = wb.create_sheet("시편_시리즈")
    cols = ["편", "시편코드", "권", "유형", "제작 구분", "본편 A코드", "제목", "첫 3초 훅", "핵심 구절",
            "한 줄 메시지", "저자·표제", "장소·배경", "교리·연출 주의", "회차 표기", "시편 롱폼",
            "유튜브 제목안", "업로드 예정일", "요일", "다음 편 예고", "대본", "교리감수", "영상", "업로드", "비고"]
    header(ws, 1, cols, [5, 8, 7, 10, 18, 9, 20, 30, 10, 32, 24, 22, 32, 10, 8, 40, 11, 5, 26, 8, 8, 8, 8, 16])
    REUSE = PatternFill("solid", fgColor="E2EFDA")
    for i, ps in enumerate(psalms, 2):
        m = ps["reuse"]
        vals = [ps["no"], f"P{ps['no']:03d}", ps["book"], ps["kind"],
                f"재사용 (본편 {m['book']['testament']} #{m['track_no']:03d})" if m else "신규 제작",
                m["code"] if m else "", ps["title"], ps["hook"], ps["kv"], ps["msg"], ps["author"],
                ps["scene"], ps["caution"]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=14, value=f"=\"시편 #\"&TEXT(A{i},\"000\")")
        ws.cell(row=i, column=15, value=ps["long"])
        ws.cell(row=i, column=16, value=f"=\"[\"&{P['시리즈 이름(가칭)']}&\" \"&N{i}&\"] \"&G{i}&\" | 시편 \"&A{i}&\"편\"")
        ws.cell(row=i, column=17, value=(
            f"={P['시편 시리즈 시작일 (월요일)']}+7*INT((A{i}-1)/3)+CHOOSE(MOD(A{i}-1,3)+1,{DAYS[0]},{DAYS[1]},{DAYS[2]})"))
        ws.cell(row=i, column=17).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=18, value=f"=CHOOSE(WEEKDAY(Q{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\")")
        ws.cell(row=i, column=19, value=f"=\"다음 편: \"&G{i + 1}" if i < n else "시편 시리즈 완결")
        for j in range(20, 24):
            ws.cell(row=i, column=j, value="미착수")
    body(ws, 2, n, len(cols), center_cols=(1, 2, 3, 4, 5, 6, 9, 14, 15, 17, 18, 20, 21, 22, 23))
    for i, ps in enumerate(psalms, 2):
        if ps["reuse"]:
            for j in (5, 6):
                ws.cell(row=i, column=j).fill = REUSE
        if ps["caution"]:
            ws.cell(row=i, column=13).font = f(False, 10, "C00000")
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
    dv = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    dv.add(f"T2:W{n}")
    ws.add_data_validation(dv)
    ws.cell(row=1, column=5).comment = Comment("재사용: 본편 영상을 '시편 #' 제목으로 재업로드(초록) / 신규: 새로 제작", "제작표")
    ws.cell(row=1, column=17).comment = Comment("주 3편(화·목·일). 시작일은 안내 시트에서 변경", "제작표")

    wl = wb.create_sheet("시편_롱폼")
    lcols = ["순번", "롱폼코드", "롱폼 제목", "권", "본문 범위", "편수", "그중 신규", "예상 길이(초)",
             "예상 길이(분)", "공개 예정일", "요일", "대본", "교리감수", "영상", "업로드", "비고"]
    header(wl, 1, lcols, [6, 9, 30, 12, 16, 6, 8, 9, 9, 11, 5, 8, 8, 8, 8, 16])
    for i, (a, b, title) in enumerate(LONGS, 2):
        books = sorted({x["book"] for x in psalms if a <= x["no"] <= b})
        vals = [i - 1, f"PL{i - 1:02d}", title, "·".join(books), f"시편 {a}~{b}편"]
        for j, v in enumerate(vals, 1):
            wl.cell(row=i, column=j, value=v)
        rngO = f"'시편_시리즈'!$O$2:$O${n}"
        wl.cell(row=i, column=6, value=f"=COUNTIF({rngO},B{i})")
        wl.cell(row=i, column=7, value=f"=COUNTIFS({rngO},B{i},'시편_시리즈'!$E$2:$E${n},\"신규 제작\")")
        wl.cell(row=i, column=8, value=f"=F{i}*{P['롱폼에 넣는 A 1편 길이(초)']}+{P['롱폼 인트로+아웃트로+연결(초)']}")
        wl.cell(row=i, column=9, value=f"=ROUND(H{i}/60,1)")
        wl.cell(row=i, column=10, value=f"=_xlfn.MAXIFS('시편_시리즈'!$Q$2:$Q${n},{rngO},B{i})+1")
        wl.cell(row=i, column=10).number_format = "yyyy-mm-dd"
        wl.cell(row=i, column=11, value=f"=CHOOSE(WEEKDAY(J{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\")")
        for j in range(12, 16):
            wl.cell(row=i, column=j, value="미착수")
    nl = len(LONGS) + 1
    body(wl, 2, nl, len(lcols), center_cols=(1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15))
    dvl = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    dvl.add(f"L2:O{nl}")
    wl.add_data_validation(dvl)
    return n, nl


def check_items(n, nl, A):
    """누락점검 시트에 넣을 (항목, 현재 값, 기대 값)."""
    rngA = f"'시편_시리즈'!$A$2:$A${n}"
    return [
        ("시편 시리즈 편수", f"=COUNT({rngA})", 150),
        ("시편 1~150편이 한 번씩 (번호 합계)", f"=SUM({rngA})", 150 * 151 // 2),
        ("시편 번호 중복", f"=SUMPRODUCT((COUNTIF({rngA},{rngA})>1)*1)", 0),
        ("시편 재사용 편수 = 본편 시편 편수", f"=COUNTIF('시편_시리즈'!$F$2:$F${n},\"A*\")", f"=COUNTIF({A('I')},\"시편\")"),
        ("시편 롱폼에 배정되지 않은 편", f"=COUNT({rngA})-SUMPRODUCT(COUNTIF('시편_시리즈'!$O$2:$O${n},'시편_롱폼'!$B$2:$B${nl}))", 0),
    ]
