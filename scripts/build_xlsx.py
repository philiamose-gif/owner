"""성경 애니메이션 제작표(엑셀) 생성 스크립트.

data/*.txt 의 분할 데이터를 읽어 제작표.xlsx 를 만든다.
실행: python3 scripts/build_xlsx.py
"""
import glob
import os
import re
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "성경애니메이션_제작표.xlsx")

CHAPTERS = {
    "창세기": 50, "출애굽기": 40, "레위기": 27, "민수기": 36, "신명기": 34,
    "여호수아": 24, "사사기": 21, "룻기": 4, "사무엘상": 31, "사무엘하": 24,
    "열왕기상": 22, "열왕기하": 25, "역대상": 29, "역대하": 36, "에스라": 10,
    "느헤미야": 13, "에스더": 10, "욥기": 42, "시편": 150, "잠언": 31,
    "전도서": 12, "아가": 8, "이사야": 66, "예레미야": 52, "예레미야애가": 5,
    "에스겔": 48, "다니엘": 12, "호세아": 14, "요엘": 3, "아모스": 9,
    "오바댜": 1, "요나": 4, "미가": 7, "나훔": 3, "하박국": 3, "스바냐": 3,
    "학개": 2, "스가랴": 14, "말라기": 4,
    "마태복음": 28, "마가복음": 16, "누가복음": 24, "요한복음": 21,
    "사도행전": 28, "로마서": 16, "고린도전서": 16, "고린도후서": 13,
    "갈라디아서": 6, "에베소서": 6, "빌립보서": 4, "골로새서": 4,
    "데살로니가전서": 5, "데살로니가후서": 3, "디모데전서": 6, "디모데후서": 4,
    "디도서": 3, "빌레몬서": 1, "히브리서": 13, "야고보서": 5, "베드로전서": 5,
    "베드로후서": 3, "요한일서": 5, "요한이서": 1, "요한삼서": 1, "유다서": 1,
    "요한계시록": 22,
}

# 전체를 연속으로 다루지 않고 핵심을 골라 다루는 책
SELECTIVE = {
    "시편": "150편 중 대표 시편 31편을 5권 구조 순서로 선별",
    "마가복음": "마태복음과 겹치는 사건은 제외, 마가 고유 장면 위주 선별",
    "누가복음": "탄생 이야기와 누가 고유 비유·장면 위주 선별",
    "요한복음": "21장 전체를 다루되 일부 강화(講話)는 대표 단락으로 압축",
    "역대상": "사무엘서와 겹치는 내용은 압축, 역대기 고유 관점 중심",
    "역대하": "열왕기와 겹치는 내용은 압축, 역대기 고유 관점 중심",
}

FONT = "Arial"
C_HEAD = "1F3864"
C_SUB = "D9E1F2"
C_INPUT = "FFF2CC"
C_OT = "F2F2F2"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(bold=False, size=10, color="000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)


WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------------------------------------------------------------- 데이터 읽기
def load():
    books, bs, rows = [], [], []
    book = None
    cur_b = None
    for path in sorted(glob.glob(os.path.join(ROOT, "data", "*.txt"))):
        for raw in open(path, encoding="utf-8"):
            line = raw.rstrip("\n")
            if not line.strip():
                continue
            if line.startswith("#"):
                parts = line[1:].strip().split("|")
                if len(parts) == 5 and ":" not in parts[0]:
                    name, abbr, testament, genre, default = [p.strip() for p in parts]
                    book = dict(name=name, abbr=abbr, testament=testament,
                                genre=genre, default=default)
                    books.append(book)
                continue
            if line.startswith("B "):
                p = line[2:].split("|")
                if len(p) not in (4, 5):
                    sys.exit(f"B 형식 오류: {line}")
                cur_b = dict(code=f"B{len(bs) + 1:03d}", title=p[0], intro=p[1],
                             christ=p[2], doctrine=p[3],
                             note=p[4] if len(p) == 5 else "", books=[])
                bs.append(cur_b)
                continue
            if line.startswith("A"):
                thematic = line.startswith("A* ")
                p = line[3:].split("|") if thematic else line[2:].split("|")
                if len(p) < 7:
                    sys.exit(f"A 형식 오류: {line}")
                p += [""] * (8 - len(p))
                rng, title, hook, kv, msg, people, place, caution = [x.strip() for x in p[:8]]
                if book["name"] not in cur_b["books"]:
                    cur_b["books"].append(book["name"])
                if book["name"] == "시편":
                    disp = f"시편 {rng}편"
                else:
                    disp = f"{book['abbr']} {rng}"
                rows.append(dict(
                    code=f"A{len(rows) + 1:04d}", b=cur_b, book=book, rng=rng,
                    disp=disp, kind="주제" if thematic else book["default"],
                    title=title, hook=hook, kv=f"{book['abbr']} {kv}", msg=msg,
                    people=people, place=place, caution=caution))
                continue
            sys.exit(f"알 수 없는 줄: {line}")
    return books, bs, rows


def chapters_of(rng):
    """'1:26-2:3' -> {1,2},  '5:1-32' -> {5},  '23' -> {23}"""
    if ":" not in rng:
        return {int(rng)}
    a, _, b = rng.partition("-")
    c1 = int(a.split(":")[0])
    c2 = int(b.split(":")[0]) if ":" in b else c1
    return set(range(c1, c2 + 1))


def span_text(first, lastr):
    """B타입 본문 범위: 첫 A의 시작 ~ 마지막 A의 끝 (예: 창 1:1~3:24)."""
    def ends(r):
        a, _, b = r["rng"].partition("-")
        if not b:
            return a, a
        if ":" in a and ":" not in b:
            b = a.split(":")[0] + ":" + b
        return a, b
    start, end = ends(first)[0], ends(lastr)[1]
    if first["book"]["name"] == "시편":
        start = f"시편 {start}"
        end = f"{end}편" if lastr["book"] is first["book"] else end
        if lastr["book"] is first["book"]:
            return f"{start}~{end}"
    else:
        start = f"{first['book']['abbr']} {start}"
    if lastr["book"] is not first["book"]:
        end = f"{lastr['book']['abbr']} {end}"
    return f"{start}~{end}"


def check(books, bs, rows):
    problems = []
    for bk in books:
        if bk["name"] not in CHAPTERS:
            problems.append(f"장 수 미등록: {bk['name']}")
            continue
        covered = set()
        for r in rows:
            if r["book"] is bk:
                covered |= chapters_of(r["rng"])
        missing = sorted(set(range(1, CHAPTERS[bk["name"]] + 1)) - covered)
        extra = sorted(c for c in covered if c > CHAPTERS[bk["name"]])
        if extra:
            problems.append(f"{bk['name']}: 존재하지 않는 장 {extra}")
        if missing and bk["name"] not in SELECTIVE:
            problems.append(f"{bk['name']}: 빠진 장 {missing}")
    if len(books) != 66:
        problems.append(f"책 수 {len(books)}권 (66권이어야 함)")
    for b in bs:
        n = sum(1 for r in rows if r["b"] is b)
        if n == 0:
            problems.append(f"{b['code']} 비어 있음")
    return problems


# ---------------------------------------------------------------- 공통 서식
def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = f(True, 10, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = CENTER
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def body(ws, r1, r2, c2, center_cols=()):
    for row in ws.iter_rows(min_row=r1, max_row=r2, max_col=c2):
        for c in row:
            c.font = f()
            c.border = BORDER
            c.alignment = CENTER if c.column in center_cols else WRAP


def title_block(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(True, 16, C_HEAD)
    ws["A2"] = sub
    ws["A2"].font = f(False, 10, "595959")


def section(ws, row, text, ncol=6):
    c = ws.cell(row=row, column=1, value=text)
    c.font = f(True, 12, "FFFFFF")
    for i in range(1, ncol + 1):
        ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor=C_HEAD)
    return row + 1


def table(ws, row, titles, data, ncol=None, wrap_rows=True):
    ncol = ncol or len(titles)
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.border = BORDER
        c.alignment = CENTER
    for d in data:
        row += 1
        for i, v in enumerate(d, 1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = f()
            c.border = BORDER
            c.alignment = WRAP
    return row + 2


# ---------------------------------------------------------------- 시트 작성
STATUS = '"미착수,진행중,완료,보류"'


def build(books, bs, rows, groups):
    wb = Workbook()
    last = len(rows) + 1  # A타입 시트의 마지막 데이터 행
    n_long = len(groups) + 1
    n_b = len(bs) + 1

    def A(col):  # A타입 시트의 열 전체 범위
        return f"'A타입_제작표'!${col}$2:${col}${last}"

    # ===== 1. 안내 =====
    ws = wb.active
    ws.title = "안내"
    title_block(ws, "성경 애니메이션 시리즈 제작표 (대한예수교장로회 합동측 기준)",
                f"창세기~요한계시록 66권 장면·주제 분할 · 시즌제 | 작성일 {date.today():%Y-%m-%d} | 개역개정 표기 기준")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 66

    r = section(ws, 4, "제작 변수 (노란 칸만 수정하세요)", 4)
    s1o = sum(1 for x in rows if x["season"] == 1 and x["book"]["testament"] == "구약")
    s1n = sum(1 for x in rows if x["season"] == 1 and x["book"]["testament"] == "신약")
    params = [
        ("시리즈 이름(가칭)", "1분 성경", "참고 영상 분석 후 확정. 유튜브 제목안에 자동 반영"),
        ("숏폼 1편 길이(초)", 60, "A타입 업로드본 길이. 권장 45~60초"),
        ("롱폼에 넣는 A 1편 길이(초)", 52, "롱폼으로 묶을 때 다음 편 예고(약 8초)를 뺀 길이"),
        ("롱폼 인트로+아웃트로+연결(초)", 120, "진행 캐릭터 인트로 30초 + 중간 질문 10초 + 아웃트로 60초 + 연결 20초"),
        ("시즌1 시작일 (월요일)", date(2027, 1, 4), "예시 날짜입니다. 반드시 월요일로 입력하세요 (신약 월·수·금·토 계산 기준)"),
        ("시즌2 시작일 (월요일)", None, "기본값: 시즌1이 끝난 뒤 1주 쉬고 시작. 원하는 월요일 날짜를 직접 입력해도 됩니다"),
        ("시즌1 목표: 구약 편수", s1o, "6개월(26주) × 매일 1편"),
        ("시즌1 목표: 신약 편수", s1n, "6개월(26주) × 주 4편"),
    ]
    prow = {}
    for label, val, note in params:
        ws.cell(row=r, column=1, value=label).font = f(True)
        c = ws.cell(row=r, column=3, value=val)
        c.fill = PatternFill("solid", fgColor=C_INPUT)
        c.font = f(False, 10, "0000FF")
        c.border = BORDER
        if "시작일" in label:
            c.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4, value=note).font = f(False, 9, "595959")
        prow[label] = r
        r += 1
    P = {k: f"'안내'!$C${v}" for k, v in prow.items()}
    ws.cell(row=prow["시즌2 시작일 (월요일)"], column=3).value = (
        f"={P['시즌1 시작일 (월요일)']}+7*(ROUNDUP(MAX(COUNTIFS({A('C')},1,{A('D')},\"구약\")/7,"
        f"COUNTIFS({A('C')},1,{A('D')},\"신약\")/4),0)+1)")

    r = section(ws, r + 1, "업로드 규칙", 4)
    for label, text in [
        ("구약 트랙", "매일 1편 (월~일)"),
        ("신약 트랙", "주 4편 (월·수·금·토)"),
        ("시즌 1 (약 6개월)", "구원 역사의 큰 줄기 + 누구나 아는 이야기 위주로 고른 편을 성경 순서대로 업로드"),
        ("시즌 2", "시즌 1에서 다루지 않은 나머지 편을 성경 순서대로 업로드. 시즌 1 + 시즌 2 = 전체(빠짐없음, '누락점검' 시트에서 확인)"),
        ("회차 번호", "트랙별로 이어서 매김 (예: 시즌1 구약 #001~#182 → 시즌2 구약 #183~). 성경 순서는 A코드로 유지"),
        ("롱폼", "시즌 중에는 '롱폼_시즌별'(시즌·트랙별 7편 이상 묶음)을 올리고, 전체 완료 후 '롱폼_통합본'(성경 순서 84편)을 정주행·교회 교육용으로 사용"),
    ]:
        ws.cell(row=r, column=1, value=label).font = f(True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(row=r, column=2, value=text)
        c.font = f()
        c.alignment = WRAP
        ws.row_dimensions[r].height = 30
        r += 1

    r = section(ws, r + 1, "전체 규모 (자동 계산)", 4)
    stats = [
        ("숏폼 총 편수", f"=COUNTA({A('B')})", "편", None),
        ("시즌1 구약 / 신약", f"=COUNTIFS({A('C')},1,{A('D')},\"구약\")&\" / \"&COUNTIFS({A('C')},1,{A('D')},\"신약\")", "편", None),
        ("시즌2 구약 / 신약", f"=COUNTIFS({A('C')},2,{A('D')},\"구약\")&\" / \"&COUNTIFS({A('C')},2,{A('D')},\"신약\")", "편", None),
        ("시즌1 종료일", f"=_xlfn.MAXIFS({A('U')},{A('C')},1)", "", "date"),
        ("시즌2 종료일 (전체 완주)", f"=_xlfn.MAXIFS({A('U')},{A('C')},2)", "", "date"),
        ("시즌1 기간", f"=ROUND((C{{r4}}-{P['시즌1 시작일 (월요일)']}+1)/7,1)", "주", None),
        ("전체 기간", f"=ROUND((C{{r5}}-{P['시즌1 시작일 (월요일)']}+1)/7,1)", "주 (쉬는 주 포함)", None),
        ("시즌별 롱폼 편수", f"=COUNTA('롱폼_시즌별'!$B$2:$B${n_long})", "편", None),
        ("통합본 롱폼 편수", f"=COUNTA('롱폼_통합본'!$B$2:$B${n_b})", "편", None),
        ("숏폼 총 분량", f"=ROUND(C{{r0}}*{P['숏폼 1편 길이(초)']}/3600,1)", "시간", None),
    ]
    r0 = r
    for k, (label, formula, unit, fmt) in enumerate(stats):
        formula = formula.replace("{r0}", str(r0)).replace("{r4}", str(r0 + 3)).replace("{r5}", str(r0 + 4))
        ws.cell(row=r, column=1, value=label).font = f(True)
        c = ws.cell(row=r, column=3, value=formula)
        c.font = f()
        c.border = BORDER
        if fmt == "date":
            c.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4, value=unit).font = f(False, 9, "595959")
        r += 1

    r = section(ws, r + 1, "분할 기준", 4)
    rules = [
        ("A타입 (숏폼·릴스)", "45~60초, 장면 1개 · 메시지 1개. 이야기 책은 '장면형', 시가서·서신서·예언 등은 '주제형'으로 분할"),
        ("롱폼", "A타입 7편 이상을 하나의 이야기 단위로 묶어 약 8분 이상"),
        ("진행 캐릭터 (절충안)", "A타입: 마지막 2~3초 등장(다음 편 안내) / 롱폼: 인트로·중간 질문·아웃트로 진행"),
        ("스토리형 원칙", "쉬운 말 내레이션 + 핵심 구절 1개만 개역개정 인용. 성경에 없는 대사·장면은 최소화하고 '상상 장면'으로 감수"),
        ("구속사 관점", "모든 롱폼에 '그리스도 연결점'을 두어 도덕 교훈에 그치지 않고 그리스도를 가리키도록 구성"),
        ("시즌1 연결 내레이션", "시즌1은 중간 편을 건너뛰므로, A타입 '연결 내레이션' 칸에 적힌 건너뛴 내용을 1~2문장으로 요약해 흐름을 이어 줌"),
        ("선별 분할한 책", "; ".join(f"{k}: {v}" for k, v in SELECTIVE.items())),
    ]
    for label, text in rules:
        ws.cell(row=r, column=1, value=label).font = f(True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(row=r, column=2, value=text)
        c.font = f()
        c.alignment = WRAP
        ws.row_dimensions[r].height = 30 if len(text) < 90 else 60
        r += 1

    r = section(ws, r + 1, "시트 구성", 4)
    for name, desc in [
        ("A타입_제작표", "숏폼 전체 798편: 시즌, 트랙 회차, 본문, 제목, 첫 3초 훅, 핵심 구절, 메시지, 인물, 배경, 교리·연출 주의, 연결 내레이션, 업로드일, 제작 상태"),
        ("롱폼_시즌별", "시즌 중 올릴 롱폼: 시즌·트랙별 묶음, 공개 예정일"),
        ("롱폼_통합본", "성경 순서 롱폼 84편: 진행자 인트로 질문, 그리스도 연결점, 핵심 교리"),
        ("누락점검", "시즌1+시즌2=전체, 중복·미배정·회차 연속성·권별 장 점검 (모두 '정상'이어야 함)"),
        ("권별요약", "66권별 편수(시즌별), 분량, 분할 방식"),
        ("인물_캐릭터시트", "주요 인물 외형·의상 가안과 삼위 하나님·천사·사탄 표현 규칙, 진행 캐릭터(가안)"),
        ("교리감수_기준", "합동측 신앙 표준, 핵심 교리 체크, 이단 해석 경계 목록, 감수 절차"),
        ("제작템플릿", "A·롱폼 구성 공식, 프롬프트 5블록, 기술 규격, 메타데이터, 저작권·업로드 체크리스트"),
        ("절기특집", "교회 절기별 추천 편성"),
        ("표기통일", "개역개정 기준 용어·인명 표기"),
    ]:
        ws.cell(row=r, column=1, value=name).font = f(True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(row=r, column=2, value=desc)
        c.font = f()
        c.alignment = WRAP
        r += 1
    ws.cell(row=r + 1, column=1, value="※ 노란 칸 = 수정 가능한 입력칸 / 상태 열은 목록에서 선택 / 그 외 숫자는 수식으로 자동 계산").font = f(False, 9, "C00000")
    ws.cell(row=r + 2, column=1, value="※ 시즌 배정을 바꾸려면 data/시즌1_선정.lst 를 고친 뒤 scripts/build_xlsx.py 로 다시 만드세요 (누락 점검이 함께 실행됨)").font = f(False, 9, "595959")

    # ===== 2. A타입_제작표 =====
    wa = wb.create_sheet("A타입_제작표")
    a_cols = ["순번", "A코드", "시즌", "트랙", "트랙 회차", "회차 표기", "통합본 B", "시즌 롱폼",
              "책", "본문", "유형", "제목", "첫 3초 훅", "핵심 구절", "한 줄 메시지", "등장인물",
              "장소·배경", "교리·연출 주의", "연결 내레이션", "유튜브 제목안", "업로드 예정일",
              "요일", "다음 편 예고", "대본", "교리감수", "영상", "업로드", "비고",
              "트랙키", "시즌 내 순서"]
    widths = [6, 8, 5, 6, 7, 10, 7, 8, 10, 16, 6, 20, 30, 12, 32, 22, 22, 34, 40, 40, 11,
              5, 24, 8, 8, 8, 8, 16, 9, 7]
    header(wa, 1, a_cols, widths)
    for i, rw in enumerate(rows, 2):
        vals = [i - 1, rw["code"], rw["season"], rw["book"]["testament"]]
        for j, v in enumerate(vals, 1):
            wa.cell(row=i, column=j, value=v)
        wa.cell(row=i, column=5, value=(
            f"=IF(C{i}=1,COUNTIFS($C$2:C{i},1,$D$2:D{i},D{i}),"
            f"COUNTIFS($C$2:$C${last},1,$D$2:$D${last},D{i})+COUNTIFS($C$2:C{i},2,$D$2:D{i},D{i}))"))
        wa.cell(row=i, column=6, value=f"=D{i}&\" #\"&TEXT(E{i},\"000\")")
        rest = [rw["b"]["code"], rw["long"], rw["book"]["name"], rw["disp"], rw["kind"], rw["title"],
                rw["hook"], rw["kv"], rw["msg"], rw["people"], rw["place"], rw["caution"], rw["bridge"]]
        for j, v in enumerate(rest, 7):
            wa.cell(row=i, column=j, value=v)
        wa.cell(row=i, column=20, value=f"=\"[\"&{P['시리즈 이름(가칭)']}&\" \"&F{i}&\"] \"&L{i}&\" | \"&J{i}")
        start = f"IF(C{i}=1,{P['시즌1 시작일 (월요일)']},{P['시즌2 시작일 (월요일)']})"
        wa.cell(row=i, column=21, value=(
            f"=IF(D{i}=\"구약\",{start}+AD{i}-1,"
            f"{start}+7*INT((AD{i}-1)/4)+CHOOSE(MOD(AD{i}-1,4)+1,0,2,4,5))"))
        wa.cell(row=i, column=21).number_format = "yyyy-mm-dd"
        wa.cell(row=i, column=22, value=f"=CHOOSE(WEEKDAY(U{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\")")
        nx = f"MATCH(D{i}&\"#\"&(E{i}+1),$AC$2:$AC${last},0)"
        wa.cell(row=i, column=23, value=(
            f"=IFERROR(IF(INDEX($C$2:$C${last},{nx})=C{i},\"다음 편: \",\"시즌1 완결 · 시즌2 첫 편: \")"
            f"&INDEX($L$2:$L${last},{nx}),\"{rw['book']['testament']} 트랙 완결\")"))
        for j in range(24, 28):
            wa.cell(row=i, column=j, value="미착수")
        wa.cell(row=i, column=29, value=f"=D{i}&\"#\"&E{i}")
        wa.cell(row=i, column=30, value=f"=COUNTIFS($C$2:C{i},C{i},$D$2:D{i},D{i})")
    body(wa, 2, last, len(a_cols), center_cols=(1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 14, 21, 22, 24, 25, 26, 27, 29, 30))
    S1_FILL = PatternFill("solid", fgColor="E2EFDA")
    for i, rw in enumerate(rows, 2):
        if rw["season"] == 1:
            for j in (3, 4, 5, 6):
                wa.cell(row=i, column=j).fill = S1_FILL
        if rw["book"]["testament"] == "구약":
            for j in (1, 2, 7, 8, 9, 10):
                wa.cell(row=i, column=j).fill = PatternFill("solid", fgColor=C_OT)
        if rw["caution"]:
            wa.cell(row=i, column=18).font = f(False, 10, "C00000")
        if rw["season"] == 1 and rw.get("gap"):
            wa.cell(row=i, column=19).font = f(False, 10, "1F3864")
    wa.freeze_panes = "M2"
    wa.auto_filter.ref = f"A1:{get_column_letter(len(a_cols))}{last}"
    dv = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    dv.add(f"X2:AA{last}")
    wa.add_data_validation(dv)
    wa.cell(row=1, column=3).comment = Comment("1 = 시즌1(6개월, 초록), 2 = 시즌2. 바꾸려면 data/시즌1_선정.lst 수정 후 재생성", "제작표")
    wa.cell(row=1, column=5).comment = Comment("트랙(구약/신약)별 업로드 순번. 시즌1을 먼저 매기고 시즌2가 이어 받음", "제작표")
    wa.cell(row=1, column=19).comment = Comment("시즌1: 건너뛴 시즌2 편 요약 → 내레이션 1~2문장으로 연결 / 시즌2: 시즌1의 어느 편 다음 이야기인지", "제작표")
    wa.cell(row=1, column=29).comment = Comment("다음 편 찾기용 보조 열", "제작표")

    # ===== 3. 롱폼_시즌별 =====
    wl = wb.create_sheet("롱폼_시즌별")
    l_cols = ["순번", "롱폼코드", "시즌", "트랙", "롱폼 제목", "포함 통합본 B", "본문 범위", "A 편수",
              "예상 길이(초)", "예상 길이(분)", "첫 회차", "마지막 회차", "진행자 인트로 질문(삶의 고민)",
              "그리스도 연결점", "핵심 교리", "공개 예정일", "요일", "대본", "교리감수", "영상", "업로드", "비고"]
    header(wl, 1, l_cols, [6, 9, 5, 6, 34, 16, 26, 7, 9, 9, 11, 11, 32, 44, 32, 11, 5, 8, 8, 8, 8, 16])
    for i, g in enumerate(groups, 2):
        first, lastr = g["rows"][0], g["rows"][-1]
        titles = [b["title"] for b in g["bs"]]
        name = titles[0] if len(titles) == 1 else f"{titles[0]} ~ {titles[-1]}"
        vals = [i - 1, g["code"], g["season"], g["track"], name,
                ", ".join(b["code"] for b in g["bs"]), span_text(first, lastr)]
        for j, v in enumerate(vals, 1):
            wl.cell(row=i, column=j, value=v)
        wl.cell(row=i, column=8, value=f"=COUNTIF({A('H')},B{i})")
        wl.cell(row=i, column=9, value=f"=H{i}*{P['롱폼에 넣는 A 1편 길이(초)']}+{P['롱폼 인트로+아웃트로+연결(초)']}")
        wl.cell(row=i, column=10, value=f"=ROUND(I{i}/60,1)")
        wl.cell(row=i, column=11, value=f"=D{i}&\" #\"&TEXT(_xlfn.MINIFS({A('E')},{A('H')},B{i}),\"000\")")
        wl.cell(row=i, column=12, value=f"=D{i}&\" #\"&TEXT(_xlfn.MAXIFS({A('E')},{A('H')},B{i}),\"000\")")
        wl.cell(row=i, column=13, value=g["bs"][0]["intro"])
        wl.cell(row=i, column=14, value="\n".join(b["christ"] for b in g["bs"]))
        wl.cell(row=i, column=15, value=" / ".join(b["doctrine"] for b in g["bs"][:2]))
        wl.cell(row=i, column=16, value=f"=_xlfn.MAXIFS({A('U')},{A('H')},B{i})+1")
        wl.cell(row=i, column=16).number_format = "yyyy-mm-dd"
        wl.cell(row=i, column=17, value=f"=CHOOSE(WEEKDAY(P{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\")")
        for j in range(18, 22):
            wl.cell(row=i, column=j, value="미착수")
    body(wl, 2, n_long, len(l_cols), center_cols=(1, 2, 3, 4, 8, 9, 10, 11, 12, 16, 17, 18, 19, 20, 21))
    for i, g in enumerate(groups, 2):
        if g["season"] == 1:
            for j in (2, 3, 4):
                wl.cell(row=i, column=j).fill = S1_FILL
    wl.freeze_panes = "F2"
    wl.auto_filter.ref = f"A1:{get_column_letter(len(l_cols))}{n_long}"
    dvl = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    dvl.add(f"R2:U{n_long}")
    wl.add_data_validation(dvl)
    wl.cell(row=1, column=16).comment = Comment("포함된 마지막 숏폼 공개 다음 날", "제작표")
    wl.cell(row=1, column=7).comment = Comment("시즌1 롱폼은 건너뛴 편이 있으므로 연결 내레이션으로 이어 붙이세요", "제작표")

    # ===== 4. 롱폼_통합본 =====
    wbs = wb.create_sheet("롱폼_통합본")
    b_cols = ["순번", "B코드", "롱폼 제목", "포함 책", "본문 범위", "A 편수", "그중 시즌1",
              "예상 길이(초)", "예상 길이(분)", "첫 A코드", "마지막 A코드", "진행자 인트로 질문(삶의 고민)",
              "그리스도 연결점", "핵심 교리", "전체 주의", "모든 편 공개 완료일",
              "대본", "교리감수", "영상", "업로드", "비고"]
    header(wbs, 1, b_cols, [6, 7, 26, 18, 26, 7, 8, 9, 9, 9, 9, 32, 40, 32, 30, 12, 8, 8, 8, 8, 16])
    for i, b in enumerate(bs, 2):
        mine = [x for x in rows if x["b"] is b]
        vals = [i - 1, b["code"], b["title"], " · ".join(b["books"]), span_text(mine[0], mine[-1])]
        for j, v in enumerate(vals, 1):
            wbs.cell(row=i, column=j, value=v)
        wbs.cell(row=i, column=6, value=f"=COUNTIF({A('G')},B{i})")
        wbs.cell(row=i, column=7, value=f"=COUNTIFS({A('G')},B{i},{A('C')},1)")
        wbs.cell(row=i, column=8, value=f"=F{i}*{P['롱폼에 넣는 A 1편 길이(초)']}+{P['롱폼 인트로+아웃트로+연결(초)']}")
        wbs.cell(row=i, column=9, value=f"=ROUND(H{i}/60,1)")
        wbs.cell(row=i, column=10, value=f"=INDEX({A('B')},MATCH(B{i},{A('G')},0))")
        wbs.cell(row=i, column=11, value=f"=INDEX({A('B')},MATCH(B{i},{A('G')},0)+F{i}-1)")
        for j, v in zip((12, 13, 14, 15), (b["intro"], b["christ"], b["doctrine"], b["note"])):
            wbs.cell(row=i, column=j, value=v)
        wbs.cell(row=i, column=16, value=f"=_xlfn.MAXIFS({A('U')},{A('G')},B{i})")
        wbs.cell(row=i, column=16).number_format = "yyyy-mm-dd"
        for j in range(17, 21):
            wbs.cell(row=i, column=j, value="미착수")
    body(wbs, 2, n_b, len(b_cols), center_cols=(1, 2, 6, 7, 8, 9, 10, 11, 16, 17, 18, 19, 20))
    wbs.freeze_panes = "D2"
    wbs.auto_filter.ref = f"A1:{get_column_letter(len(b_cols))}{n_b}"
    dvb = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    dvb.add(f"Q2:T{n_b}")
    wbs.add_data_validation(dvb)
    wbs.cell(row=1, column=16).comment = Comment("시즌1·2 숏폼이 모두 공개된 날. 통합본은 이 날 이후 정주행·교회 교육용으로 공개", "제작표")

    # ===== 5. 누락점검 =====
    checks_sheet(wb, books, rows, groups, last, n_long, A, P)

    # ===== 6. 권별요약 =====
    wk = wb.create_sheet("권별요약")
    k_cols = ["순서", "책", "약어", "구약/신약", "장르", "장 수", "A 편수", "시즌1", "시즌2",
              "통합본 B 수", "숏폼 분량(분)", "성경 순서 시작", "성경 순서 끝", "분할 방식"]
    header(wk, 1, k_cols, [6, 14, 6, 9, 10, 7, 8, 7, 7, 10, 10, 10, 10, 50])
    for i, bk in enumerate(books, 2):
        vals = [i - 1, bk["name"], bk["abbr"], bk["testament"], bk["genre"], CHAPTERS[bk["name"]]]
        for j, v in enumerate(vals, 1):
            wk.cell(row=i, column=j, value=v)
        wk.cell(row=i, column=7, value=f"=COUNTIF({A('I')},B{i})")
        wk.cell(row=i, column=8, value=f"=COUNTIFS({A('I')},B{i},{A('C')},1)")
        wk.cell(row=i, column=9, value=f"=COUNTIFS({A('I')},B{i},{A('C')},2)")
        wk.cell(row=i, column=10, value=f"=SUMPRODUCT(({A('I')}=B{i})/COUNTIFS({A('I')},{A('I')},{A('G')},{A('G')}))")
        wk.cell(row=i, column=11, value=f"=ROUND(G{i}*{P['숏폼 1편 길이(초)']}/60,1)")
        wk.cell(row=i, column=12, value=f"=_xlfn.MINIFS({A('A')},{A('I')},B{i})")
        wk.cell(row=i, column=13, value=f"=_xlfn.MAXIFS({A('A')},{A('I')},B{i})")
        wk.cell(row=i, column=14, value=SELECTIVE.get(bk["name"], "전체 장을 순서대로 빠짐없이 분할"))
    tot = len(books) + 2
    wk.cell(row=tot, column=2, value="합계")
    for col in "FGHIK":
        idx = ord(col) - 64
        wk.cell(row=tot, column=idx, value=f"=SUM({col}2:{col}{tot - 1})")
    body(wk, 2, tot, len(k_cols), center_cols=tuple(range(1, 14)))
    for j in range(1, len(k_cols) + 1):
        wk.cell(row=tot, column=j).font = f(True)
        wk.cell(row=tot, column=j).fill = PatternFill("solid", fgColor=C_SUB)
    wk.freeze_panes = "C2"

    # ===== 7~11 참고 시트 =====
    from sheets_ref import characters, doctrine, templates, seasons, terms
    characters(wb, rows, last, f, header, body, title_block, section, table, C_INPUT, BORDER, WRAP)
    doctrine(wb, f, title_block, section, table)
    templates(wb, f, title_block, section, table)
    seasons(wb, rows, f, header, body, title_block)
    terms(wb, f, title_block, section, table)

    wb.save(OUT)
    return OUT


def checks_sheet(wb, books, rows, groups, last, n_long, A, P):
    ws = wb.create_sheet("누락점검")
    title_block(ws, "누락 점검", "시즌1 + 시즌2 = 성경 전체가 되는지 확인합니다. 결과가 모두 '정상'이어야 합니다.")
    for c, w in zip("ABCDEFGHI", (40, 14, 14, 12, 12, 10, 10, 12, 40)):
        ws.column_dimensions[c].width = w
    r = section(ws, 4, "1. 전체 점검 (자동 계산)", 5)
    for j, t in enumerate(["점검 항목", "현재 값", "기대 값", "결과"], 1):
        c = ws.cell(row=r, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.border = BORDER
    n = len(rows)
    s1 = sum(1 for x in rows if x["season"] == 1)
    items = [
        ("숏폼 총 편수", f"=COUNTA({A('B')})", n),
        ("시즌1 편수", f"=COUNTIF({A('C')},1)", s1),
        ("시즌2 편수", f"=COUNTIF({A('C')},2)", n - s1),
        ("시즌1 + 시즌2 (전체와 같아야 함)", f"=COUNTIF({A('C')},1)+COUNTIF({A('C')},2)", f"=COUNTA({A('B')})"),
        ("시즌이 지정되지 않은 편", f"=COUNTA({A('B')})-COUNTIF({A('C')},1)-COUNTIF({A('C')},2)", 0),
        ("A코드 중복", f"=SUMPRODUCT((COUNTIF({A('B')},{A('B')})>1)*1)", 0),
        ("트랙 회차 중복", f"=SUMPRODUCT((COUNTIF({A('AC')},{A('AC')})>1)*1)", 0),
        ("구약 트랙 마지막 회차 = 구약 편수", f"=_xlfn.MAXIFS({A('E')},{A('D')},\"구약\")", f"=COUNTIF({A('D')},\"구약\")"),
        ("신약 트랙 마지막 회차 = 신약 편수", f"=_xlfn.MAXIFS({A('E')},{A('D')},\"신약\")", f"=COUNTIF({A('D')},\"신약\")"),
        ("시즌1 구약 편수 = 목표", f"=COUNTIFS({A('C')},1,{A('D')},\"구약\")", f"={P['시즌1 목표: 구약 편수']}"),
        ("시즌1 신약 편수 = 목표", f"=COUNTIFS({A('C')},1,{A('D')},\"신약\")", f"={P['시즌1 목표: 신약 편수']}"),
        ("같은 트랙·같은 날 겹친 업로드", f"=SUMPRODUCT((COUNTIFS({A('D')},{A('D')},{A('U')},{A('U')})>1)*1)", 0),
        ("시즌 롱폼에 배정되지 않은 편", f"=COUNTA({A('B')})-SUMPRODUCT(COUNTIF({A('H')},'롱폼_시즌별'!$B$2:$B${n_long}))", 0),
        ("통합본 B에 배정되지 않은 편", f"=COUNTA({A('B')})-SUMPRODUCT(COUNTIF({A('G')},'롱폼_통합본'!$B$2:$B${len({id(x['b']) for x in rows}) + 1}))", 0),
        ("시즌2 시작일이 시즌1 종료일보다 뒤인가 (1 = 예)", f"=IF({P['시즌2 시작일 (월요일)']}>_xlfn.MAXIFS({A('U')},{A('C')},1),1,0)", 1),
        ("시작일이 월요일인가 (시즌1·2 모두, 2 = 예)", f"=(WEEKDAY({P['시즌1 시작일 (월요일)']},2)=1)+(WEEKDAY({P['시즌2 시작일 (월요일)']},2)=1)", 2),
    ]
    r += 1
    first = r
    for label, cur, exp in items:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=cur)
        ws.cell(row=r, column=3, value=exp)
        ws.cell(row=r, column=4, value=f"=IF(B{r}=C{r},\"정상\",\"확인 필요\")")
        for j in range(1, 5):
            ws.cell(row=r, column=j).font = f(j == 4)
            ws.cell(row=r, column=j).border = BORDER
        r += 1
    ws.cell(row=r, column=1, value="전체 결과").font = f(True)
    c = ws.cell(row=r, column=4, value=f"=IF(COUNTIF(D{first}:D{r - 1},\"정상\")={len(items)},\"모두 정상\",\"확인 필요 있음\")")
    c.font = f(True, 11, "C00000")
    c.border = BORDER
    r += 3

    r = section(ws, r, "2. 권별 점검", 9)
    cols = ["책", "장 수", "다룬 장 수", "A 편수", "시즌1", "시즌2", "시즌 합계", "결과", "빠진 장 / 비고"]
    for j, t in enumerate(cols, 1):
        c = ws.cell(row=r, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    for bk in books:
        covered = set()
        for x in rows:
            if x["book"] is bk:
                covered |= chapters_of(x["rng"])
        total = CHAPTERS[bk["name"]]
        missing = sorted(set(range(1, total + 1)) - covered)
        if not missing:
            note = "없음 (모든 장 포함)"
        elif bk["name"] == "시편":
            note = f"대표 31편 선별 — 나머지 {len(missing)}편은 이번 시리즈 범위 밖 (의도적)"
        else:
            note = f"{', '.join(map(str, missing))}장 — {SELECTIVE[bk['name']]} (의도적)"
        ws.cell(row=r, column=1, value=bk["name"])
        ws.cell(row=r, column=2, value=total)
        ws.cell(row=r, column=3, value=len(covered & set(range(1, total + 1))))
        ws.cell(row=r, column=4, value=f"=COUNTIF({A('I')},A{r})")
        ws.cell(row=r, column=5, value=f"=COUNTIFS({A('I')},A{r},{A('C')},1)")
        ws.cell(row=r, column=6, value=f"=COUNTIFS({A('I')},A{r},{A('C')},2)")
        ws.cell(row=r, column=7, value=f"=E{r}+F{r}")
        ws.cell(row=r, column=8, value=f"=IF(G{r}=D{r},\"정상\",\"확인 필요\")")
        ws.cell(row=r, column=9, value=note)
        for j in range(1, 10):
            cc = ws.cell(row=r, column=j)
            cc.font = f(j == 8, 10, "C00000" if (j == 9 and missing) else "000000")
            cc.border = BORDER
            cc.alignment = WRAP if j == 9 else Alignment(horizontal="center")
        r += 1
    r += 2

    r = section(ws, r, "3. 시즌1 흐름 점검 (건너뛴 구간)", 9)
    for j, t in enumerate(["트랙", "시즌1 편수", "바로 이어지는 편", "건너뛴 구간이 있는 편", "가장 긴 건너뜀"], 1):
        c = ws.cell(row=r, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.border = BORDER
    r += 1
    for t in ("구약", "신약"):
        s1rows = [x for x in rows if x["season"] == 1 and x["book"]["testament"] == t]
        gaps = [x for x in s1rows if x["gap"]]
        worst = max(s1rows, key=lambda x: x["gap"])
        vals = [t, len(s1rows), len(s1rows) - len(gaps), len(gaps),
                f"{worst['gap']}편 ({t} #{worst['track_no']:03d} '{worst['title']}' 앞)"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = f()
            c.border = BORDER
        r += 1
    ws.cell(row=r + 1, column=1, value="※ 건너뛴 구간이 있는 편은 A타입_제작표 '연결 내레이션' 칸에 건너뛴 편 목록이 적혀 있습니다. 대본 작성 시 1~2문장으로 이어 주세요.").font = f(False, 9, "595959")
    ws.cell(row=r + 2, column=1, value="※ 이 시트 2번의 '다룬 장 수'와 '빠진 장'은 제작표를 만들 때 스크립트가 본문 범위를 장 단위로 검사한 결과입니다.").font = f(False, 9, "595959")


if __name__ == "__main__":
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from season import assign, long_groups, bridges
    books, bs, rows = load()
    problems = check(books, bs, rows)
    print(f"책 {len(books)}권 / 통합본 B {len(bs)}편 / A {len(rows)}편")
    for p in problems:
        print("확인 필요:", p)
    if any("빠진 장" in p or "존재하지" in p or "66권" in p for p in problems):
        sys.exit(1)
    assign(rows)
    groups = long_groups(rows, bs)
    bridges(rows)
    s1 = [x for x in rows if x["season"] == 1]
    assert len(s1) + sum(1 for x in rows if x["season"] == 2) == len(rows)
    assert all("long" in x for x in rows), "시즌 롱폼 미배정 편 있음"
    sizes = sorted(len(g["rows"]) for g in groups)
    print(f"시즌1 {len(s1)}편 / 시즌2 {len(rows) - len(s1)}편 / 시즌 롱폼 {len(groups)}편 (편수 {sizes[0]}~{sizes[-1]})")
    print("저장:", build(books, bs, rows, groups))
