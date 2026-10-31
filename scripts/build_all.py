"""두 프로젝트 제작표를 한 번에 만들고 검사한다.

순서: 역본 대조 → 프로젝트1 엑셀 → 프로젝트2 엑셀 → 수식 계산(LibreOffice) → 누락점검 결과 요약
실행: python3 scripts/build_all.py
"""
import glob
import json
import os
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUTS = [
    ("프로젝트1", os.path.join(ROOT, "scripts", "build_xlsx.py"), os.path.join(ROOT, "성경애니메이션_제작표.xlsx")),
    ("프로젝트2", os.path.join(ROOT, "project2", "build.py"), os.path.join(ROOT, "project2", "5분성경읽기_제작표.xlsx")),
]


def recalc_script():
    """xlsx 스킬의 수식 계산 스크립트 위치 (없으면 계산 단계를 건너뜀)."""
    hits = glob.glob("/root/.claude/skills/**/xlsx/scripts/recalc.py", recursive=True)
    return hits[0] if hits else None


def run(cmd):
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
        sys.exit(f"실패: {' '.join(cmd)}")
    return r.stdout


def summary(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["누락점검"]
    bad, total = [], None
    for row in ws.iter_rows(min_row=1, max_col=4, values_only=True):
        if row[0] == "전체 결과":
            total = row[3]
        elif row[3] == "확인 필요":
            bad.append(row[0])
    return total, bad


if __name__ == "__main__":
    print(run([sys.executable, os.path.join(ROOT, "scripts", "verify_bible.py")]).strip().splitlines()[-1])
    recalc = recalc_script()
    for name, script, out in OUTPUTS:
        print(f"[{name}]", run([sys.executable, script]).strip().splitlines()[-1])
        if recalc:
            res = json.loads(run([sys.executable, recalc, out, "300"]))
            if res.get("status") != "success":
                sys.exit(f"{name} 수식 오류: {res}")
            total, bad = summary(out)
            print(f"[{name}] 수식 {res['total_formulas']}개 오류 0 · 누락점검: {total}", *(f"\n  확인 필요: {b}" for b in bad))
        else:
            print(f"[{name}] 수식 계산 도구가 없어 엑셀에서 열 때 계산됩니다")
