"""프로젝트2 '5분 성경읽기' 제작표 생성.

개역한글 성경 전체(창세기~요한계시록)를 성인이 자연스럽게 소리 내어 읽어 5분 안팎이 되도록 끊고,
구약·신약 두 트랙으로 6개월(182일) 동안 하루 여러 편씩 올리는 일정을 만든다. 각 편의 TTS 대본도 함께 만든다.

실행: python3 project2/build.py          (제작표)
      python3 project2/build.py --tts    (제작표 + project2/tts/ 에 편별 TTS 대본 텍스트)
"""
import math
import os
import re
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))
import build_xlsx as bx  # noqa: E402  (서식 도구, 66권 정보 재사용)
import verify_bible as vb  # noqa: E402  (역본 본문)

OUT = os.path.join(HERE, "5분성경읽기_제작표.xlsx")
TTS_DIR = os.path.join(HERE, "tts")
TARGET_DAYS = 182  # 6개월

# 낭독 길이 기준 — 바꾸면 다시 실행해야 분할이 바뀜 (엑셀의 속도 칸은 예상 시간만 다시 계산)
SPM = 300        # 1분에 읽는 음절 수 (성인이 차분하고 자연스럽게 소리 내어 읽는 속도)
CHAPTER_SEC = 2  # 장 안내("창세기 1장")와 장 사이 쉼
TARGET, LO, HI = 300, 210, 375  # 목표 5분, 허용 3분 30초~6분 15초
SENTENCE_END = "라다요냐까"      # 개역한글은 마침표가 없어 절 끝 어미로 문장 끝을 판단

f, header, body, section, table = bx.f, bx.header, bx.body, bx.section, bx.table
C_INPUT, C_SUB, BORDER = bx.C_INPUT, bx.C_SUB, bx.BORDER
STATUS = bx.STATUS
SLOTS = ["06:00", "08:00", "12:00", "15:00", "19:00", "21:00"]  # 하루 업로드 시간 (17시 전후는 프로젝트1)

# 장르 → 화면 배경 (사람 얼굴·하나님 형상 없이 풍경과 빛, 본문 자막이 잘 보이게 어둡고 단순하게)
BACKGROUNDS = {
    "율법서": "광야와 장막, 돌판, 저녁 하늘",
    "역사서": "옛 성읍과 성벽, 언덕길, 강가",
    "시가서": "새벽빛, 잔잔한 물, 들판의 바람",
    "선지서": "바위산, 구름 사이 빛줄기, 두루마리",
    "복음서": "갈릴리 호수, 올리브 나무, 들길",
    "서신서": "등불 아래 펼친 두루마리, 창가의 햇빛",
    "예언서": "별이 가득한 하늘, 빛나는 성의 실루엣",
}
GENRE_BG = {"율법서": "율법서", "역사서": "역사서", "시가서": "시가서", "선지서": "선지서", "복음서": "복음서",
            "예언서": "예언서", "바울서신": "서신서", "옥중서신": "서신서", "목회서신": "서신서", "일반서신": "서신서"}


def syllables(t):
    return len(re.findall(r"[가-힣0-9]", t))


def chapter_word(book):
    return "편" if book["name"] == "시편" else "장"


def split_book(book):
    """한 권을 5분 안팎으로 끊는다. 끊는 곳은 장 끝을 가장 선호하고, 장 안에서는 문장이 끝나는 절에서만."""
    ab = book["abbr"]
    units = []  # (장, 절, 음절, 끊을 때 벌점)
    keys = [k for k in vb.KRV if k[0] == ab]
    ch_syl = {}
    for _, c, v in keys:
        ch_syl[c] = ch_syl.get(c, 0) + syllables(vb.KRV[(ab, c, v)][0])
    long_ch = lambda c: ch_syl[c] / SPM * 60 > (HI if book["name"] == "시편" else TARGET)
    for i, (_, c, v) in enumerate(keys):
        t, fl = vb.KRV[(ab, c, v)]
        silent = fl.startswith("merged") or vb.PLACEHOLDER.fullmatch(t)
        syl = 0 if silent else syllables(t)
        nxt = keys[i + 1] if i + 1 < len(keys) else None
        if nxt is None or nxt[1] != c:
            pen = 0
        elif not silent and t.rstrip()[-1:] in SENTENCE_END:
            pen = 3 if long_ch(c) else 40 if book["name"] == "시편" else 12  # 짧은 장·시편 한 편은 되도록 통째로
        else:
            pen = 40
        units.append((c, v, syl, pen))
    n = len(units)
    pre = [0]
    for u in units:
        pre.append(pre[-1] + u[2])

    def dur(a, b):  # units[a:b] 낭독 초
        chaps = len({units[k][0] for k in range(a, b) if k == a or units[k][1] == 1})
        return (pre[b] - pre[a]) / SPM * 60 + chaps * CHAPTER_SEC

    if dur(0, n) <= HI * 1.2:  # 7분 30초 이하인 짧은 책은 나누지 않고 한 편
        return [(units[0][:2], units[-1][:2], dur(0, n), 0)]
    best = [0.0] + [math.inf] * n
    back = [0] * (n + 1)
    for b in range(1, n + 1):
        for a in range(b - 1, -1, -1):
            d = dur(a, b)
            if d > HI * 1.6:
                break
            if best[a] == math.inf:
                continue
            cost = ((d - TARGET) / 30) ** 2 + units[b - 1][3]
            if d < LO or d > HI:
                cost += 200
            if best[a] + cost < best[b]:
                best[b], back[b] = best[a] + cost, a
    cuts, b = [], n
    while b:
        cuts.append((back[b], b))
        b = back[b]
    return [(units[a][:2], units[b - 1][:2], dur(a, b), units[b - 1][3]) for a, b in reversed(cuts)]


def fmt_ref(c1, v1, c2, v2):
    if c1 == c2:
        return f"{c1}:{v1}-{v2}" if v1 != v2 else f"{c1}:{v1}"
    return f"{c1}:{v1}-{c2}:{v2}"


def disp_ko(book, c1, v1, c2, v2):
    """'창세기 1장 1절 ~ 2장 3절' 식 표기 (장 전체면 절 생략)"""
    w = chapter_word(book)
    whole_start, whole_end = v1 == 1, v2 == vb.last_verse(book["abbr"], c2)
    if whole_start and whole_end:
        return f"{book['name']} {c1}{w}" if c1 == c2 else f"{book['name']} {c1}~{c2}{w}"
    a = f"{c1}{w} {v1}절"
    b = f"{v2}절" if c1 == c2 else f"{c2}{w} {v2}절"
    return f"{book['name']} {a} ~ {b}"


def tts_script(book, c1, v1, c2, v2):
    """낭독용 대본: 장이 시작할 때 '창세기 1장' 안내, 절 번호는 읽지 않음, 시편 표제는 본문이라 읽음."""
    ab, lines, para = book["abbr"], [], []
    for (_, c, v) in vb.verses_in(ab, fmt_ref(c1, v1, c2, v2)):
        if v == 1 or not lines and not para:
            if para:
                lines.append(" ".join(para))
                para = []
            head = f"{book['name']} {c}{chapter_word(book)}" + ("" if v == 1 else f" {v}절부터")
            lines.append(head + ".")
        t, fl = vb.KRV[(ab, c, v)]
        if fl.startswith("merged") or vb.PLACEHOLDER.fullmatch(t):
            continue
        para.append(pause_marks(re.sub(r"\{([^}]*)\}\s*", r"\1. ", t).strip()))  # 시편 표제 뒤 쉼
    if para:
        lines.append(" ".join(para))
    return "\n".join(lines)


FINAL = re.compile(r"(니라|느니라|더라|도다|로다|지어다|소서|니이다|나이다|노라|리라|이니라)( )(?!하시|하고|하며|하매|하니|하였|하더|하여|한즉|함|말하|이르|가로|대답)")


def pause_marks(t):
    """TTS가 쉬어 읽도록 문장 끝에 마침표, 이어지는 절 끝에 쉼표만 덧붙인다 (낱말은 바꾸지 않음, 자막은 원문 그대로)."""
    t = FINAL.sub(r"\1.\2", t)
    if t[-1:] in ".,?!)":
        return t
    return t + ("." if t[-1:] in SENTENCE_END else ",")


def spots_text(book, ref):
    """원본과 대조할 자리만: '1:14 일 자와 / 1:30 식물로 주노라' (현재 띄어쓰기 그대로 표시)"""
    out = []
    for k in vb.verses_in(book["abbr"], ref):
        fl = vb.KRV[k][1]
        if fl.startswith("break:"):
            out += [f"{k[1]}:{k[2]} {sp}" for sp in fl[6:].split("/")]
    return f"띄어쓰기 {len(out)}곳: " + " / ".join(out) if out else ""


def load():
    books = bx.load()[0]
    segs = []
    for track in ("구약", "신약"):
        no = 0
        for b in (x for x in books if x["testament"] == track):
            for (c1, v1), (c2, v2), d, pen in split_book(b):
                no += 1
                ref = fmt_ref(c1, v1, c2, v2)
                segs.append(dict(track=track, no=no, book=b, ref=ref, disp=f"{b['abbr']} {ref}",
                                 title=disp_ko(b, c1, v1, c2, v2), sec=round(d),
                                 cut="장 끝" if pen == 0 else "문장 끝",
                                 chapters=c2 - c1 + 1, verses=len(vb.verses_in(b["abbr"], ref)),
                                 tts=tts_script(b, c1, v1, c2, v2), syl=sum(
                                     0 if fl.startswith("merged") or vb.PLACEHOLDER.fullmatch(t) else syllables(t)
                                     for t, fl in (vb.KRV[k] for k in vb.verses_in(b["abbr"], ref))),
                                 whole_book=(c1, v1) == (1, 1) and c2 == bx.CHAPTERS[b["name"]]
                                 and v2 == vb.last_verse(b["abbr"], c2)))
    # 일정: 트랙마다 182일에 고르게 나누고, 하루 안에서는 구약 → 신약 순서로 시간대를 채움
    for track in ("구약", "신약"):
        tr = [s for s in segs if s["track"] == track]
        for k, s in enumerate(tr):
            s["day"] = k * TARGET_DAYS // len(tr) + 1
    slot = {}
    for s in sorted(segs, key=lambda s: (s["day"], s["track"] != "구약", s["no"])):
        slot[s["day"]] = slot.get(s["day"], 0) + 1
        s["slot"] = slot[s["day"]]
    return books, segs


def check(books, segs):
    problems = vb.check_project2(segs)
    if max(s["slot"] for s in segs) > len(SLOTS):
        problems.append(f"하루 편수가 업로드 시간대({len(SLOTS)}개)보다 많음")
    for s in segs:
        if not s["whole_book"] and not LO <= s["sec"] <= HI:
            problems.append(f"{s['title']} 낭독 {s['sec']}초 (허용 범위 밖)")
    return problems


def build(books, segs):
    wb = Workbook()
    n = len(segs)
    last = n + 1
    R = lambda c: f"'읽기목록'!${c}$2:${c}${last}"
    total_min = sum(s["sec"] for s in segs) / 60

    # ===== 안내 =====
    ws = wb.active
    ws.title = "안내"
    bx.title_block(ws, "프로젝트2 · 5분 성경읽기 제작표",
                   f"개역한글 성경 전체를 5분 안팎 {n}편으로 · 6개월(182일) · 하루 여러 편 | 작성일 {date.today():%Y-%m-%d}")
    for c, w in zip("ABCD", (30, 18, 16, 64)):
        ws.column_dimensions[c].width = w
    r = section(ws, 4, "제작 변수 (노란 칸만 수정하세요)", 4)
    params = [
        ("시리즈 이름(가칭)", "5분 성경읽기", "같은 채널 안의 재생목록 이름. 제목·해시태그에 반영"),
        ("시작일", date(2027, 1, 1), "예시 날짜입니다. 프로젝트1과 상관없이 정하면 됩니다"),
        ("낭독 속도 (음절/분)", SPM, "예상 낭독 시간 계산용. 분할 자체를 바꾸려면 build.py의 SPM을 고친 뒤 다시 실행"),
        ("장 안내·쉼 (초/장)", CHAPTER_SEC, "'창세기 1장' 안내와 장 사이 쉼"),
        ("인트로 (초)", 5, "시그니처 소리 + '5분 성경읽기, 창세기 1장' 화면"),
        ("마무리 (초)", 5, "출처(개역한글) 표시 + 다음 편 안내 화면"),
    ] + [(f"업로드 시간 {k}", t, "하루 안의 순서대로 사용 (구약 먼저, 신약 나중)") for k, t in enumerate(SLOTS, 1)]
    prow = {}
    for label, val, note in params:
        ws.cell(row=r, column=1, value=label).font = f(True)
        c = ws.cell(row=r, column=3, value=val)
        c.fill = PatternFill("solid", fgColor=C_INPUT)
        c.font = f(False, 10, "0000FF")
        c.border = BORDER
        if isinstance(val, date):
            c.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4, value=note).font = f(False, 9, "595959")
        prow[label] = r
        r += 1
    P = {k: f"'안내'!$C${v}" for k, v in prow.items()}
    slot_range = f"'안내'!$C${prow['업로드 시간 1']}:$C${prow[f'업로드 시간 {len(SLOTS)}']}"

    r = section(ws, r + 1, "전체 규모 (자동 계산)", 4)
    for label, formula, unit, fmt in [
        ("전체 편수", f"=COUNTA({R('B')})", "편", None),
        ("구약 / 신약", f"=COUNTIF({R('C')},\"구약\")&\" / \"&COUNTIF({R('C')},\"신약\")", "편", None),
        ("성경 전체 절 수", f"=SUM({R('L')})", f"절 / 개역한글 {len(vb.KRV)}절", None),
        ("전체 낭독 시간", f"=ROUND(SUM({R('O')})/3600,1)", "시간 (인트로·마무리 제외)", None),
        ("하루 평균 편수", f"=ROUND(COUNTA({R('B')})/(MAX({R('E')})),1)", "편", None),
        ("하루 최대 편수", f"=MAX({R('H')})", "편", None),
        ("마지막 업로드일", f"=MAX({R('F')})", "", "date"),
        ("전체 기간", f"=MAX({R('F')})-{P['시작일']}+1", "일 (목표 182일)", None),
        ("띄어쓰기 대조가 필요한 편", f"=COUNTIF({R('V')},\"*대조 필요\")", "편 (줄바꿈 띄어쓰기를 정하지 못한 절 포함)", None),
    ]:
        ws.cell(row=r, column=1, value=label).font = f(True)
        c = ws.cell(row=r, column=3, value=formula)
        c.border = BORDER
        if fmt == "date":
            c.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4, value=unit).font = f(False, 9, "595959")
        r += 1

    r = section(ws, r + 1, "운영 원칙", 4)
    for label, text in [
        ("성경 읽기에만 집중", "해설·진행자 없이 개역한글 본문 낭독(TTS)과 자막만. 설명이 필요하면 유튜브 설명란·인스타 게시글에만"),
        ("분량", f"성인이 자연스럽게 소리 내어 읽는 속도(1분 {SPM}음절 안팎)로 5분 안팎(3분 30초~6분 15초). 장이 끝나는 곳에서 먼저 끊고, 장 안에서는 문장이 끝나는 절에서만 끊음. 7분 30초 이하인 짧은 책은 나누지 않고 한 편"),
        ("트랙", "구약(창세기→말라기)과 신약(마태복음→요한계시록)을 동시에 진행해 182일째 함께 끝남. 하루 구약 3~4편 + 신약 1~2편"),
        ("채널·형식", "프로젝트1과 같은 채널의 '5분 성경읽기' 재생목록. 5분 영상은 쇼츠·릴스 길이(최대 3분)를 넘으므로 유튜브 일반 영상(16:9)이 기본"),
        ("성경 본문", "개역한글판만. 본문과 TTS 대본은 reference/krv.tsv(개역한글 원문 변환본)에서 자동 생성. '대조 필요' 편은 띄어쓰기를 원본 PDF와 확인"),
        ("일정", "'일자별' 시트가 그날 올릴 편 목록. '오늘의 성경읽기'로 하루치를 묶어 안내 가능"),
    ]:
        ws.cell(row=r, column=1, value=label).font = f(True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(row=r, column=2, value=text)
        c.font = f()
        c.alignment = bx.WRAP
        ws.row_dimensions[r].height = 30
        r += 1

    # ===== 읽기목록 =====
    wr = wb.create_sheet("읽기목록")
    cols = ["순번", "코드", "트랙", "트랙 회차", "일차", "업로드일", "요일", "하루 순서", "업로드 시간", "책", "본문 범위",
            "절 수", "장 수", "음절 수", "예상 낭독(초)", "예상 낭독", "영상 길이", "끊은 곳", "제목안", "해시태그",
            "설명란 문구", "본문 상태", "대조할 곳", "화면 배경", "TTS 대본 (개역한글)", "TTS 음성", "영상", "업로드",
            "조회수", "평균 시청률(%)", "비고"]
    header(wr, 1, cols, [6, 8, 6, 7, 6, 11, 5, 6, 8, 10, 30, 6, 6, 7, 8, 8, 8, 8, 44, 40, 44, 16, 22, 26, 60,
                         8, 8, 8, 8, 9, 16])
    for i, s in enumerate(segs, 2):
        b = s["book"]
        _, flags = vb.text(b["abbr"], s["ref"], strip_heading=False)
        status = "자동 입력 · 대조 필요" if flags & vb.REVIEW else "자동 입력"
        vals = {1: i - 1, 2: f"R{i - 1:04d}", 3: s["track"], 4: s["no"], 5: s["day"], 8: s["slot"], 10: b["name"],
                11: s["title"], 12: s["verses"], 13: s["chapters"],
                14: s["syl"],
                18: "책 전체" if s["whole_book"] else s["cut"], 22: status,
                23: spots_text(b, s["ref"]) or ", ".join(sorted(flags - {vb.F_HEADING})), 24: f"{GENRE_BG[b['genre']]}: {BACKGROUNDS[GENRE_BG[b['genre']]]}",
                25: s["tts"], 26: "미착수", 27: "미착수", 28: "미착수"}
        for j, val in vals.items():
            wr.cell(row=i, column=j, value=val)
        wr.cell(row=i, column=6, value=f"={P['시작일']}+E{i}-1").number_format = "yyyy-mm-dd"
        wr.cell(row=i, column=7, value=f"=CHOOSE(WEEKDAY(F{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\")")
        wr.cell(row=i, column=9, value=f"=INDEX({slot_range},H{i})")
        wr.cell(row=i, column=15, value=f"=ROUND(N{i}/{P['낭독 속도 (음절/분)']}*60+M{i}*{P['장 안내·쉼 (초/장)']},0)")
        wr.cell(row=i, column=16, value=f"=TEXT(O{i}/86400,\"m:ss\")")
        wr.cell(row=i, column=17, value=f"=TEXT((O{i}+{P['인트로 (초)']}+{P['마무리 (초)']})/86400,\"m:ss\")")
        wr.cell(row=i, column=19, value=f"={P['시리즈 이름(가칭)']}&\" | {s['title']}\"")
        wr.cell(row=i, column=20, value=(
            f"=\"#\"&SUBSTITUTE({P['시리즈 이름(가칭)']},\" \",\"\")&\" #성경읽기 #성경통독 #{b['name']} #개역한글\""))
        wr.cell(row=i, column=21, value=f"{s['title']} (개역한글)\n성경 본문: 개역한글판 · 낭독: AI 음성")
    body(wr, 2, last, len(cols), center_cols=(1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 12, 13, 14, 15, 16, 17, 18, 22, 26, 27, 28, 29, 30))
    for i, s in enumerate(segs, 2):
        for j in (29, 30):
            wr.cell(row=i, column=j).fill = PatternFill("solid", fgColor=C_INPUT)
        if s["track"] == "구약":
            for j in (1, 2, 3, 4):
                wr.cell(row=i, column=j).fill = PatternFill("solid", fgColor=bx.C_OT)
        wr.row_dimensions[i].height = 30
    wr.freeze_panes = "K2"
    wr.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    dv1 = DataValidation(type="list", formula1='"자동 입력,자동 입력 · 대조 필요,대조 완료"', allow_blank=True)
    dv1.add(f"V2:V{last}")
    dv2 = DataValidation(type="list", formula1=STATUS, allow_blank=True)
    dv2.add(f"Z2:AB{last}")
    wr.add_data_validation(dv1)
    wr.add_data_validation(dv2)

    # ===== 일자별 =====
    wd = wb.create_sheet("일자별")
    header(wd, 1, ["일차", "날짜", "요일", "편수", "구약", "신약", "낭독 합계(분)"], [6, 11, 5, 6, 60, 40, 10])
    for d in range(1, TARGET_DAYS + 1):
        i = d + 1
        today = [s for s in segs if s["day"] == d]
        rng = lambda t: " / ".join(s["title"] for s in today if s["track"] == t)
        wd.cell(row=i, column=1, value=d)
        wd.cell(row=i, column=2, value=f"={P['시작일']}+A{i}-1").number_format = "yyyy-mm-dd"
        wd.cell(row=i, column=3, value=f"=CHOOSE(WEEKDAY(B{i},2),\"월\",\"화\",\"수\",\"목\",\"금\",\"토\",\"일\")")
        wd.cell(row=i, column=4, value=f"=COUNTIF({R('E')},A{i})")
        wd.cell(row=i, column=5, value=rng("구약"))
        wd.cell(row=i, column=6, value=rng("신약"))
        wd.cell(row=i, column=7, value=f"=ROUND(SUMIF({R('E')},A{i},{R('O')})/60,0)")
    body(wd, 2, TARGET_DAYS + 1, 7, center_cols=(1, 2, 3, 4, 7))
    wd.freeze_panes = "B2"

    # ===== 권별점검 =====
    wk = wb.create_sheet("권별점검")
    header(wk, 1, ["책", "구약/신약", "편수", "장 수", "절 수(분할 합)", "절 수(개역한글)", "첫 업로드일", "마지막 업로드일", "결과"],
           [14, 9, 6, 7, 12, 13, 12, 14, 10])
    for i, b in enumerate(books, 2):
        wk.cell(row=i, column=1, value=b["name"])
        wk.cell(row=i, column=2, value=b["testament"])
        wk.cell(row=i, column=3, value=f"=COUNTIF({R('J')},A{i})")
        wk.cell(row=i, column=4, value=bx.CHAPTERS[b["name"]])
        wk.cell(row=i, column=5, value=f"=SUMIF({R('J')},A{i},{R('L')})")
        wk.cell(row=i, column=6, value=sum(1 for k in vb.KRV if k[0] == b["abbr"]))
        wk.cell(row=i, column=7, value=f"=_xlfn.MINIFS({R('F')},{R('J')},A{i})")
        wk.cell(row=i, column=8, value=f"=_xlfn.MAXIFS({R('F')},{R('J')},A{i})")
        for j in (7, 8):
            wk.cell(row=i, column=j).number_format = "yyyy-mm-dd"
        wk.cell(row=i, column=9, value=f"=IF(AND(C{i}>0,E{i}=F{i}),\"정상\",\"확인 필요\")")
    body(wk, 2, len(books) + 1, 9, center_cols=(2, 3, 4, 5, 6, 7, 8, 9))
    wk.freeze_panes = "B2"

    # ===== 누락점검 =====
    wc = wb.create_sheet("누락점검")
    bx.title_block(wc, "누락 점검", "결과가 모두 '정상'이어야 합니다.")
    for c, w in zip("ABCD", (48, 14, 14, 12)):
        wc.column_dimensions[c].width = w
    LOC, HIC = R("O"), R("R")
    items = [
        ("전체 편수", f"=COUNTA({R('B')})", n),
        ("성경 전체 절 수 (분할 합 = 개역한글)", f"=SUM({R('L')})", len(vb.KRV)),
        ("절 수가 맞지 않거나 빠진 책", "=COUNTIF('권별점검'!$I$2:$I$67,\"확인 필요\")", 0),
        ("포함된 책 수", "=COUNTIF('권별점검'!$C$2:$C$67,\">0\")", 66),
        (f"{LO // 60}분 {LO % 60}초보다 짧은 편 (짧은 책 한 권 전체는 제외)", f"=COUNTIFS({LOC},\"<{LO}\",{HIC},\"<>책 전체\")", 0),
        (f"{HI // 60}분 {HI % 60}초보다 긴 편 (짧은 책 한 권 전체는 제외)", f"=COUNTIFS({LOC},\">{HI}\",{HIC},\"<>책 전체\")", 0),
        ("같은 날·같은 시간 겹친 업로드", f"=SUMPRODUCT((COUNTIFS({R('E')},{R('E')},{R('H')},{R('H')})>1)*1)", 0),
        ("트랙 회차 중복", f"=SUMPRODUCT((COUNTIFS({R('C')},{R('C')},{R('D')},{R('D')})>1)*1)", 0),
        ("6개월(182일) 안에 끝나는가 (1 = 예)", f"=IF(MAX({R('F')})-{P['시작일']}+1<=182,1,0)", 1),
        ("TTS 대본이 비어 있는 편", f"=COUNTBLANK({R('Y')})", 0),
    ]
    r = section(wc, 4, "점검 항목", 4)
    for j, t in enumerate(["항목", "현재 값", "기대 값", "결과"], 1):
        c = wc.cell(row=r, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.border = BORDER
    first = r + 1
    for k, (label, cur, exp) in enumerate(items):
        rr = first + k
        wc.cell(row=rr, column=1, value=label)
        wc.cell(row=rr, column=2, value=cur)
        wc.cell(row=rr, column=3, value=exp)
        wc.cell(row=rr, column=4, value=f"=IF(B{rr}=C{rr},\"정상\",\"확인 필요\")")
        for j in range(1, 5):
            wc.cell(row=rr, column=j).border = BORDER
            wc.cell(row=rr, column=j).font = f(j == 4)
    rr = first + len(items)
    wc.cell(row=rr, column=1, value="전체 결과").font = f(True)
    wc.cell(row=rr, column=4, value=f"=IF(COUNTIF(D{first}:D{rr - 1},\"정상\")={len(items)},\"모두 정상\",\"확인 필요 있음\")").font = f(True, 11, "C00000")
    wc.cell(row=rr + 2, column=1, value="참고: 속도 칸을 바꾸면 예상 시간만 다시 계산됩니다. 길이 점검이 어긋나면 build.py의 SPM을 맞춰 다시 실행하세요.").font = f(False, 9, "595959")

    # ===== 제작템플릿 =====
    wt = wb.create_sheet("제작템플릿")
    bx.title_block(wt, "5분 성경읽기 제작 템플릿", "성경 읽기에만 집중: 해설·진행자 없이 본문 낭독(TTS)과 자막")
    for c, w in zip("ABCD", (24, 60, 40, 30)):
        wt.column_dimensions[c].width = w
    r = section(wt, 4, "1. 영상 구성 (5분 안팎)", 4)
    r = table(wt, r, ["구간", "내용", "연결 칸", "비고"], [
        ("0~5초", "시그니처 소리 + 시리즈 이름 + 오늘 본문 범위", "제목안", "프로젝트1과 다른 짧은 소리"),
        ("본문", "TTS 낭독. 화면에는 지금 읽는 절을 절 번호와 함께 2~4줄씩 자막으로", "TTS 대본", "절 번호는 화면에만 쓰고 읽지 않음"),
        ("장 바뀜", "'창세기 2장' 안내 + 1~2초 쉼, 화면에 장 제목 카드", "장 수", "장 안내·쉼 칸"),
        ("마지막 5초", "'성경 본문: 개역한글판' + 다음 편 범위 안내", "-", "재생목록 연속 재생 유도"),
    ])
    r = section(wt, r, "2. TTS 낭독 기준", 4)
    r = table(wt, r, ["항목", "기준", "비고"], [
        ("속도", f"1분 {SPM}음절 안팎 (성인이 차분하고 자연스럽게 읽는 속도). 엔진 기본 속도에서 시작해 5분짜리 한 편으로 맞춤", "예상 낭독 칸과 실제 길이를 몇 편 비교해 속도 칸 보정"),
        ("목소리", "성인 남성 또는 여성 1명을 시리즈 내내 고정. 차분하고 또렷하게", "편마다 목소리가 바뀌지 않게 같은 설정 저장"),
        ("대본", "TTS 대본 칸 그대로 사용. 장 안내('창세기 1장')와 쉬어 읽기용 쉼표·마침표만 덧붙였고 낱말은 개역한글 원문 그대로", "개역한글은 마침표가 없어 문장 끝 어미(~니라, ~더라 등)로 자동 판단. 어색한 쉼은 듣고 고침. 자막에는 쉼표·마침표를 넣지 않음"),
        ("발음 점검", "인명·지명(므두셀라, 아히도벨 등), 숫자(일백이십), 한자어 장단음을 첫 편들에서 귀로 확인하고 엔진의 발음 사전에 등록", "오독은 '비고'에 적고 다시 생성"),
        ("괄호·표제", "본문 괄호 '(그대로 되니라)'는 본문이므로 읽음. 시편 표제(다윗의 시 등)도 본문이라 읽음", "원본이 앞 절과 합쳐 인쇄한 자리 표시는 대본에서 뺐음"),
        ("엔진 선택", "한국어 신경망 음성을 제공하는 상용 TTS(클라우드 TTS, 국내 AI 성우 서비스 등) 중 선택", "유튜브 수익화·상업 이용 허용 여부와 출처 표기 조건을 약관으로 꼭 확인"),
        ("일괄 생성", "python3 project2/build.py --tts 실행 시 project2/tts/ 에 편별 대본 텍스트(R0001_….txt) 생성", "대부분의 TTS 도구에 파일 단위로 넣기 쉬움"),
    ])
    r = section(wt, r, "3. 화면", 4)
    r = table(wt, r, ["장르", "배경", "편수"],
              [(k, d, f"=COUNTIF({R('X')},\"{k}:*\")") for k, d in BACKGROUNDS.items()])
    r = table(wt, r, ["항목", "기준", "비고"], [
        ("화면 비율", "1920×1080 (16:9). 세로판이 필요하면 1080×1920으로 같은 자막 배치", "5분은 쇼츠·릴스(최대 3분) 길이를 넘음"),
        ("자막", "굵은 명조 또는 고딕, 화면 아래쪽 2~4줄, 절 번호 작게 표시", "개역한글 원문 띄어쓰기 그대로"),
        ("배경음악", "잔잔한 저작권 무료 음원, 낭독보다 충분히 작게", "찬송가·CCM 사용 금지"),
        ("배경 영상", "사람 얼굴·하나님 형상 없이 풍경과 빛. 장르별로 같은 배경 세트를 돌려 써서 제작 시간 단축", "본문 장면을 그리지 않음 (프로젝트1과 구분)"),
    ])
    r = section(wt, r, "4. 저작권·교리 원칙", 4)
    table(wt, r, ["원칙", "내용"], [
        ("본문", "개역한글판만 사용(저작재산권 보호기간 만료). 개역개정·새번역 등 다른 번역 문장을 섞지 않음"),
        ("해설", "영상에 해설을 넣지 않음. 설명란에 쓰는 글은 프로젝트1의 교리감수 기준(합동측)을 그대로 적용"),
        ("AI 음성 표기", "설명란에 '낭독: AI 음성' 표기 (플랫폼의 AI 콘텐츠 표시 정책도 확인)"),
    ])

    # ===== 성과기록 =====
    wp = wb.create_sheet("성과기록")
    bx.title_block(wp, "성과 기록", "읽기목록의 노란 칸(조회수·평균 시청률)에 입력하면 책별로 집계됩니다.")
    for j, t in enumerate(["책", "편수", "입력된 편", "평균 조회수", "평균 시청률(%)"], 1):
        c = wp.cell(row=4, column=j, value=t)
        c.font = f(True)
        c.fill = PatternFill("solid", fgColor=C_SUB)
    for k, b in enumerate(books, 5):
        wp.cell(row=k, column=1, value=b["name"])
        wp.cell(row=k, column=2, value=f"=COUNTIF({R('J')},A{k})")
        wp.cell(row=k, column=3, value=f"=COUNTIFS({R('J')},A{k},{R('AC')},\"<>\")")
        wp.cell(row=k, column=4, value=f"=IFERROR(AVERAGEIFS({R('AC')},{R('J')},A{k}),\"-\")")
        wp.cell(row=k, column=5, value=f"=IFERROR(AVERAGEIFS({R('AD')},{R('J')},A{k}),\"-\")")
    for c, w in zip("ABCDE", (14, 7, 10, 12, 14)):
        wp.column_dimensions[c].width = w

    wb.save(OUT)
    return OUT, total_min


if __name__ == "__main__":
    books, segs = load()
    problems = check(books, segs)
    secs = sorted(s["sec"] for s in segs)
    print(f"5분 읽기 {len(segs)}편 (구약 {sum(s['track'] == '구약' for s in segs)} / 신약 {sum(s['track'] == '신약' for s in segs)})"
          f" · 낭독 {secs[0]}~{secs[-1]}초, 중앙값 {secs[len(secs) // 2]}초 · 하루 최대 {max(s['slot'] for s in segs)}편")
    for p in problems:
        print("확인 필요:", p)
    if problems:
        sys.exit(1)
    if "--tts" in sys.argv:
        os.makedirs(TTS_DIR, exist_ok=True)
        for s in segs:
            name = f"R{segs.index(s) + 1:04d}_{s['title'].replace(' ', '').replace('~', '-')}.txt"
            with open(os.path.join(TTS_DIR, name), "w", encoding="utf-8") as fp:
                fp.write(s["tts"] + "\n")
        print("TTS 대본:", TTS_DIR)
    out, total = build(books, segs)
    print(f"저장: {out} (낭독 합계 {total / 60:.1f}시간)")
